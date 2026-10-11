"""Write the tailored CV, cover letter and an application pack for one job."""

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .models import Job
from .tailor import TailoredApplication

ACCENT = RGBColor(0x1F, 0x4E, 0x5F)


def slugify(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:60]


def _base_document() -> Document:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(2)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(1.5)
        section.left_margin = section.right_margin = Cm(1.8)
    return doc


def _header(doc: Document, profile: dict) -> None:
    name = doc.add_paragraph()
    run = name.add_run(profile["name"])
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = ACCENT
    contact = doc.add_paragraph(" · ".join(x for x in (profile["location"], profile["phone"], profile["email"]) if x))
    contact.paragraph_format.space_after = Pt(6)


def _heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = ACCENT
    # thin rule under the heading
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", "1F4E5F")):
        bottom.set(qn(key), value)
    border.append(bottom)
    p._p.get_or_add_pPr().append(border)


def _bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet").paragraph_format.space_after = Pt(1)


def write_cv(profile: dict, app: TailoredApplication, path: Path) -> None:
    doc = _base_document()
    _header(doc, profile)

    _heading(doc, "Profile")
    doc.add_paragraph(app.profile_statement)

    _heading(doc, "Key skills")
    doc.add_paragraph(" • ".join(app.key_skills))

    _heading(doc, "Experience")
    entries = {e["id"]: e for e in profile["experience"]}
    text_width = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    for item in app.experience:
        entry = entries[item.id]
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
        p.add_run(entry["role"]).bold = True
        where = ", ".join(x for x in (entry["employer"], entry.get("location")) if x)
        p.add_run(f" - {where}\t{entry['dates']}")
        for line in item.bullets:
            _bullet(doc, line)

    _heading(doc, "Education")
    for edu in profile["education"]:
        p = doc.add_paragraph()
        p.add_run(edu["school"]).bold = True
        p.add_run(f" - {edu['detail']}")

    if profile.get("achievements"):
        _heading(doc, "Achievements")
        for line in profile["achievements"]:
            _bullet(doc, line)

    if profile.get("interests"):
        _heading(doc, "Interests")
        for line in profile["interests"]:
            _bullet(doc, line)

    _heading(doc, "Availability")
    doc.add_paragraph(f"{profile['availability']}. Can start: {profile['start_date']}.")
    doc.save(path)


def write_cover_letter(profile: dict, app: TailoredApplication, path: Path) -> None:
    doc = _base_document()
    _header(doc, profile)
    doc.add_paragraph(date.today().strftime("%d %B %Y")).paragraph_format.space_after = Pt(12)
    for para in app.cover_letter.strip().split("\n\n"):
        doc.add_paragraph(para.strip()).paragraph_format.space_after = Pt(8)
    doc.save(path)


def write_pack(profile: dict, job: Job, app: TailoredApplication, root: Path) -> Path:
    folder = root / f"{date.today().isoformat()}-{slugify(job.employer)}-{slugify(job.title)}"
    folder.mkdir(parents=True, exist_ok=True)
    write_cv(profile, app, folder / f"CV - {profile['name']}.docx")
    write_cover_letter(profile, app, folder / f"Cover letter - {profile['name']}.docx")
    answers = "\n\n".join(f"**{a.question}**\n\n{a.answer}" for a in app.screening_answers)
    (folder / "application.md").write_text(
        f"# {job.title} - {job.employer}\n\n"
        f"- **Where:** {job.location}\n"
        f"- **Ride:** {f'~{job.ride_minutes} min by Dott ({job.ride_km} km)' if job.ride_minutes else 'unknown - check the advert'}\n"
        f"- **Pay:** {job.pay or 'not stated'}\n"
        f"- **Hours:** {job.hours or 'not stated'}\n"
        f"- **Apply:** {job.url}\n\n"
        f"> {app.fit_summary}\n\n"
        f"## Cover letter\n\n{app.cover_letter}\n\n"
        f"## Answers to common application questions\n\n{answers}\n",
        encoding="utf-8",
    )
    return folder
